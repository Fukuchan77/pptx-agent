"""Table builder for creating tables from TableBlock specifications.

This module provides functionality to add tables to PowerPoint slides
using python-pptx table capabilities.
"""

import logging

from pptx.util import Inches

from pptx_agent.pptx_wrapper.slide import SlideWrapper
from pptx_agent.schemas.visual_assets import TableBlock

logger = logging.getLogger(__name__)


def add_table_to_slide(slide: SlideWrapper, table_block: TableBlock) -> None:
    """Add a table to a slide based on TableBlock specification.

    Args:
        slide: SlideWrapper instance to add table to
        table_block: TableBlock containing table specification

    Raises:
        ValueError: If table data is invalid (empty rows/headers, mismatched columns)
    """
    # Validate table has rows
    if not table_block.rows:
        msg = "Table must have at least one row"
        raise ValueError(msg)

    # Validate table has columns (headers)
    if not table_block.headers:
        msg = "Table must have at least one column"
        raise ValueError(msg)

    # Validate all rows have the same number of columns as headers
    expected_cols = len(table_block.headers)
    for i, row in enumerate(table_block.rows):
        if len(row) != expected_cols:
            msg = f"Row {i} has {len(row)} columns, expected {expected_cols}"
            raise ValueError(msg)

    # Calculate table dimensions (rows + 1 for header row)
    num_rows = len(table_block.rows) + 1  # +1 for header row
    num_cols = len(table_block.headers)

    # Default table position and size (can be customized later)
    x = Inches(1)
    y = Inches(2)
    cx = Inches(8)
    cy = Inches(4)

    # Add table to slide using public method
    table_shape = slide.add_table(num_rows, num_cols, x, y, cx, cy)

    # Access the table object
    table = table_shape.table

    # Populate header row
    for col_idx, header_text in enumerate(table_block.headers):
        cell = table.cell(0, col_idx)
        cell.text = header_text

    # Populate data rows
    for row_idx, row_data in enumerate(table_block.rows):
        for col_idx, cell_value in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)  # +1 to skip header row
            cell.text = cell_value

    logger.info(
        "Added table to slide with %d rows and %d columns",
        num_rows,
        num_cols,
    )
